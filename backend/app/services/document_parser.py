"""
Document Specification Parser for Bharat AI Platform.

Extracts AI model, workload, cost, and latency specifications from:
- PDF documents (.pdf)
- Microsoft Word documents (.docx)
- Text and Markdown specifications (.txt, .md, .json)
"""

import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, Optional


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pypdf with fallback."""
    text_content = []
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_content.append(text)
    except Exception as e:
        # Fallback to plain string extraction if reader fails
        raw = file_bytes.decode("utf-8", errors="ignore")
        text_content.append(raw)
    return "\n".join(text_content)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from Word (.docx) using python-docx with zipfile fallback."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception:
        # Direct ZIP archive extraction fallback
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                text_nodes = root.iter()
                texts = [elem.text for elem in text_nodes if elem.text]
                return " ".join(texts)
        except Exception:
            return file_bytes.decode("utf-8", errors="ignore")


def parse_project_specification_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses a document (PDF, Word, or text/JSON) and extracts structured parameters
    for the AI Hardware Recommender.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        raw_text = extract_text_from_docx(file_bytes)
    elif ext in (".txt", ".md"):
        raw_text = file_bytes.decode("utf-8", errors="ignore")
    elif ext == ".json":
        try:
            parsed_json = json.loads(file_bytes.decode("utf-8", errors="ignore"))
            if isinstance(parsed_json, dict):
                return _normalize_json_spec(parsed_json)
        except Exception:
            raw_text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raw_text = file_bytes.decode("utf-8", errors="ignore")

    return parse_text_to_spec(raw_text, filename)


def parse_text_to_spec(text: str, source_name: str = "Project Document") -> Dict[str, Any]:
    """
    Heuristic rule engine to extract model parameters, output time / latency,
    budget, power, and workload from raw text.
    """
    text_lower = text.lower()

    # 1. Detect Model Architecture & Name
    detected_model_name = Path(source_name).stem.replace("_", " ").title()
    model_type = "CNN / Vision"
    param_billions = 0.5

    # Check for LLM patterns
    if any(k in text_lower for k in ("llama", "mistral", "qwen", "gemma", "phi-3", "phi-2", "deepseek", "chatgpt", "gpt", "falcon", "llm", "large language model")):
        model_type = "LLM"
        param_billions = 7.0
        # Check specific parameter sizes
        if "70b" in text_lower:
            param_billions = 70.0
            detected_model_name = "Llama-3-70B"
        elif "32b" in text_lower:
            param_billions = 32.0
            detected_model_name = "Qwen-2.5-32B"
        elif "14b" in text_lower:
            param_billions = 14.0
            detected_model_name = "Qwen-14B"
        elif "8b" in text_lower:
            param_billions = 8.0
            detected_model_name = "Llama-3-8B"
        elif "7b" in text_lower:
            param_billions = 7.0
            detected_model_name = "Mistral-7B"
        elif "3b" in text_lower:
            param_billions = 3.0
            detected_model_name = "Llama-3.2-3B"
        elif "1b" in text_lower or "2b" in text_lower:
            param_billions = 1.5
            detected_model_name = "Gemma-2B"
        elif "phi" in text_lower:
            param_billions = 3.8
            detected_model_name = "Phi-3-Mini"
    elif any(k in text_lower for k in ("yolo", "resnet", "mobilenet", "efficientnet", "vision", "detection", "segmentation", "cnn")):
        model_type = "CNN / Vision"
        if "yolo" in text_lower:
            detected_model_name = "YOLOv8 Object Detection"
            param_billions = 0.043  # ~43M
        elif "resnet" in text_lower:
            detected_model_name = "ResNet-50 Defect Inspector"
            param_billions = 0.025  # ~25M
        elif "mobilenet" in text_lower:
            detected_model_name = "MobileNet Edge Vision"
            param_billions = 0.0035  # ~3.5M
    elif any(k in text_lower for k in ("whisper", "speech", "audio", "voice", "asr", "wav2vec")):
        model_type = "Audio / Speech"
        detected_model_name = "Whisper ASR Speech-to-Text"
        param_billions = 1.55  # ~1.55B for large
    elif any(k in text_lower for k in ("bert", "indicbert", "embeddings", "sentiment", "token")):
        model_type = "Transformer / NLP"
        detected_model_name = "IndicBERT Multilingual Transformer"
        param_billions = 0.110  # ~110M

    # Fallback parameter search e.g. "25 million", "7 billion", "500m"
    p_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:b|billion)\s*params?", text_lower)
    if p_match:
        param_billions = float(p_match.group(1))
    else:
        m_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:m|million)\s*params?", text_lower)
        if m_match:
            param_billions = round(float(m_match.group(1)) / 1000.0, 4)

    # 2. Precision / Quantization
    precision = "FP16"
    if "int4" in text_lower or "q4" in text_lower or "awq" in text_lower or "gguf" in text_lower:
        precision = "INT4"
    elif "int8" in text_lower or "q8" in text_lower or "w8a8" in text_lower:
        precision = "INT8"
    elif "bf16" in text_lower or "bfloat16" in text_lower:
        precision = "BF16"
    elif "fp32" in text_lower:
        precision = "FP32"

    # 3. Workload
    workload = "Inference"
    if "full training" in text_lower or "pretrain" in text_lower or "train from scratch" in text_lower:
        workload = "Full Training"
    elif "lora" in text_lower or "qlora" in text_lower or "fine-tun" in text_lower or "finetun" in text_lower:
        workload = "LoRA Fine-Tuning"

    # 4. Target Environment
    target_env = "Workstation"
    if any(k in text_lower for k in ("drone", "uav", "robot", "battery", "handheld", "wearable", "ultra-low power")):
        target_env = "Drone / Battery"
    elif any(k in text_lower for k in ("edge", "cctv", "camera", "raspberry", "jetson", "iot", "gateway")):
        target_env = "Edge Appliance"
    elif any(k in text_lower for k in ("datacenter", "data center", "cloud", "cluster", "server rack", "h100", "a100", "enterprise")):
        target_env = "Datacenter"
    else:
        target_env = "Workstation"

    # 5. Output Time & Latency
    target_latency_ms = None
    lat_match = re.search(r"(?:latency|output time|response time|delay)(?:\s*(?:is|of|under|below|<=|<|:))?\s*(\d+(?:\.\d+)?)\s*(?:ms|milliseconds)", text_lower)
    if lat_match:
        target_latency_ms = float(lat_match.group(1))
    elif "real-time" in text_lower or "realtime" in text_lower:
        target_latency_ms = 30.0  # 30ms standard for real-time video/audio

    # 6. Throughput (FPS or Tokens/sec)
    target_throughput = None
    tp_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:tokens?/s|tok/s|tokens per sec|fps|frames per second)", text_lower)
    if tp_match:
        target_throughput = float(tp_match.group(1))

    # 7. Budget (INR)
    max_budget_inr = None
    b_curr_match = re.search(r"(?:₹|rs\.?|inr)\s*(\d[\d,]*)(?:\s*(?:k|thousand|lakh|lakhs|cr|crore))?", text_lower)
    if b_curr_match:
        val_str = b_curr_match.group(1).replace(",", "")
        try:
            val = float(val_str)
            context_window = text_lower[max(0, b_curr_match.start() - 10): min(len(text_lower), b_curr_match.end() + 15)]
            if "lakh" in context_window:
                val *= 100000.0
            elif "k" in context_window and val < 1000:
                val *= 1000.0
            max_budget_inr = val
        except Exception:
            pass
    if max_budget_inr is None:
        b_word_match = re.search(r"(?:budget|max cost)[^0-9\n]{0,25}(\d[\d,]+)", text_lower)
        if b_word_match:
            try:
                val = float(b_word_match.group(1).replace(",", ""))
                if val >= 100:
                    max_budget_inr = val
            except Exception:
                pass

    # 8. Power budget (Watts)
    max_power_w = None
    pow_match = re.search(r"(?:power|tdp|wattage)[^0-9\n]{0,35}(\d+(?:\.\d+)?)\s*(?:w|watts)", text_lower)
    if pow_match:
        try:
            max_power_w = float(pow_match.group(1))
        except Exception:
            pass

    # 9. Context length & Concurrency
    context_length = 4096
    ctx_match = re.search(r"(\d+)\s*(?:k)?\s*(?:tokens?|context|seq|sequence length)", text_lower)
    if ctx_match:
        c_val = int(ctx_match.group(1))
        if c_val < 256:
            c_val *= 1024  # e.g. 8k -> 8192
        context_length = min(131072, max(512, c_val))

    concurrent_users = 1
    user_match = re.search(r"(\d+)\s*(?:users?|concurrent|requests)", text_lower)
    if user_match:
        concurrent_users = min(1000, max(1, int(user_match.group(1))))

    # 10. Priority & Sovereign Indian Silicon Preference
    prefer_indian = any(k in text_lower for k in ("india", "shakti", "vega", "vicharak", "domestic", "indigenous", "atmanirbhar", "sovereign"))
    
    priority = "Balanced"
    if any(k in text_lower for k in ("cost", "cheap", "budget", "affordable", "cost-effective")):
        priority = "Cost-Efficient"
    elif any(k in text_lower for k in ("fast", "performance", "high-throughput", "low latency", "maximum speed")):
        priority = "Maximum Performance"
    elif any(k in text_lower for k in ("low power", "energy efficient", "battery", "green")):
        priority = "Low Power"
    elif prefer_indian:
        priority = "Sovereign Indian"

    summary = (
        f"Parsed specification for '{detected_model_name}': {param_billions}B parameters at {precision} for {workload} "
        f"in {target_env} environment. (Priority: {priority}"
        f"{', Indian Silicon Preferred' if prefer_indian else ''})."
    )

    return {
        "model_name": detected_model_name,
        "model_type": model_type,
        "parameter_count_b": float(param_billions),
        "precision": precision,
        "workload": workload,
        "target_environment": target_env,
        "target_latency_ms": target_latency_ms,
        "target_throughput": target_throughput,
        "max_budget_inr": max_budget_inr,
        "max_power_w": max_power_w,
        "context_length": context_length,
        "batch_size": 1,
        "concurrent_users": concurrent_users,
        "priority": priority,
        "prefer_indian": prefer_indian,
        "extracted_summary": summary,
        "raw_text_preview": text[:500] if text else "",
    }


def _normalize_json_spec(d: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "model_name": d.get("model_name") or d.get("name") or "Custom Project Model",
        "model_type": d.get("model_type") or "LLM",
        "parameter_count_b": float(d.get("parameter_count_b") or d.get("parameters_b") or 7.0),
        "precision": str(d.get("precision") or "INT4").upper(),
        "workload": d.get("workload") or "Inference",
        "target_environment": d.get("target_environment") or "Workstation",
        "target_latency_ms": float(d.get("target_latency_ms")) if d.get("target_latency_ms") else None,
        "target_throughput": float(d.get("target_throughput")) if d.get("target_throughput") else None,
        "max_budget_inr": float(d.get("max_budget_inr")) if d.get("max_budget_inr") else None,
        "max_power_w": float(d.get("max_power_w")) if d.get("max_power_w") else None,
        "context_length": int(d.get("context_length") or 4096),
        "batch_size": int(d.get("batch_size") or 1),
        "concurrent_users": int(d.get("concurrent_users") or 1),
        "priority": d.get("priority") or "Cost-Efficient",
        "prefer_indian": bool(d.get("prefer_indian", False)),
        "extracted_summary": "Parsed from uploaded JSON project specification.",
    }
